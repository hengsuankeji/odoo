from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_heading_with_color(doc, text, level=1, color=RGBColor(44, 62, 80)):
    heading = doc.add_heading(text, level)
    for run in heading.runs:
        run.font.color.rgb = color
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    return heading

def add_bullet_point(doc, text, level=0):
    p = doc.add_paragraph()
    p.style = 'List Bullet'
    p.paragraph_format.left_indent = Inches(0.5 * (level + 1))
    run = p.add_run(text)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    return p

def create_pricing_document():
    doc = Document()
    
    # 设置文档默认字体
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    
    # 添加标题
    title = doc.add_heading('恒算科技ERP系统 - 超值定价方案', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(44, 62, 80)
    
    # 产品核心优势
    add_heading_with_color(doc, '🌟 产品核心优势', 1)
    
    # 技术实力
    add_heading_with_color(doc, '技术实力', 2)
    add_bullet_point(doc, '国内软件行业TOP级开发团队')
    add_bullet_point(doc, '系统稳定性强，性能卓越')
    add_bullet_point(doc, '操作界面美观，用户体验极佳')
    
    # 功能亮点
    add_heading_with_color(doc, '功能亮点', 2)
    add_bullet_point(doc, '移动办公：随时随地查看生产状态和库存')
    add_bullet_point(doc, '智能管理：支持手机扫码，无需额外设备')
    add_bullet_point(doc, '全流程覆盖：采购、生产、销售、库存一体化')
    add_bullet_point(doc, '人力资源：员工管理、工资管理、合同管理')
    add_bullet_point(doc, '财务模块：完整的会计系统')
    add_bullet_point(doc, '无限并发：不限制同时在线人数')
    
    # 超值模块赠送
    add_heading_with_color(doc, '🎁 超值模块赠送', 1)
    
    # 基础模块
    add_heading_with_color(doc, '基础模块（全部包含）', 2)
    add_bullet_point(doc, '销售管理')
    add_bullet_point(doc, '财务管理')
    add_bullet_point(doc, '采购管理')
    add_bullet_point(doc, '库存管理')
    add_bullet_point(doc, '生产制造')
    add_bullet_point(doc, '人力资源管理')
    add_bullet_point(doc, '智能条码系统')
    
    # 赠送模块
    add_heading_with_color(doc, '赠送模块（50+个）', 2)
    add_bullet_point(doc, '客户关系管理(CRM)')
    add_bullet_point(doc, '项目管理')
    add_bullet_point(doc, '资产管理')
    add_bullet_point(doc, '质量管理')
    add_bullet_point(doc, '设备管理')
    add_bullet_point(doc, '更多模块持续更新中...')
    
    # 超值定价方案
    add_heading_with_color(doc, '💰 超值定价方案', 1)
    
    # 基础方案
    add_heading_with_color(doc, '基础方案', 2)
    add_bullet_point(doc, '软件买断价：¥188,800（原价）')
    add_bullet_point(doc, '包含所有基础模块和赠送模块')
    add_bullet_point(doc, '终身使用授权')
    
    # 年度服务费
    add_heading_with_color(doc, '年度服务费', 2)
    add_bullet_point(doc, '云服务器费用：¥9,800/年')
    add_bullet_point(doc, '技术支持维护：¥4,000/年')
    add_bullet_point(doc, '合计：¥13,800/年')
    
    # 限时优惠活动
    add_heading_with_color(doc, '🎉 限时优惠活动（2024年4-5月）', 1)
    
    # 签约优惠
    add_heading_with_color(doc, '签约优惠', 2)
    add_bullet_point(doc, '买断价直降：¥168,800（立省¥20,000）')
    add_bullet_point(doc, '首年服务费全免（价值¥13,800）')
    add_bullet_point(doc, '额外赠送3个月技术支持服务')
    add_bullet_point(doc, '免费提供系统培训（价值¥5,000）')
    
    # 老客户推荐奖励
    add_heading_with_color(doc, '老客户推荐奖励', 2)
    add_bullet_point(doc, '成功推荐新客户签约，奖励¥5,000现金')
    add_bullet_point(doc, '或等值技术支持服务')
    
    # 智能条码系统优势
    add_heading_with_color(doc, '📱 智能条码系统优势', 1)
    add_bullet_point(doc, '无需购买扫码枪，手机即可扫码')
    add_bullet_point(doc, '支持生产流水线实时扫码')
    add_bullet_point(doc, '库存管理一键扫码出入库')
    add_bullet_point(doc, '数据实时同步，操作简单便捷')
    add_bullet_point(doc, '大幅提升工作效率，降低人力成本')
    
    # 服务承诺
    add_heading_with_color(doc, '🤝 服务承诺', 1)
    add_bullet_point(doc, '7×24小时技术支持')
    add_bullet_point(doc, '免费系统升级')
    add_bullet_point(doc, '定期功能更新')
    add_bullet_point(doc, '专业培训支持')
    add_bullet_point(doc, '数据安全保障')
    
    # 页脚
    doc.add_paragraph()
    footer = doc.add_paragraph('注：本方案最终解释权归恒算科技所有')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.color.rgb = RGBColor(128, 128, 128)
    
    # 保存文档
    doc.save('contract/output/恒算科技ERP系统定价方案.docx')

if __name__ == '__main__':
    create_pricing_document() 